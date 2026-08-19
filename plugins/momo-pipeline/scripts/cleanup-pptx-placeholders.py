#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation


PROMPTS = (
    "單擊此處",
    "点击此处",
    "按一下以新增",
    "Click to add",
    "Click to edit",
)


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return shape.text_frame.text.strip()


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def should_remove(shape, remove_all_placeholders: bool) -> bool:
    text = shape_text(shape)
    if any(prompt in text for prompt in PROMPTS):
        return True
    if getattr(shape, "is_placeholder", False):
        if remove_all_placeholders:
            return True
        return text == ""
    return False


def cleanup(input_path: Path, output_path: Path, remove_all_placeholders: bool) -> tuple[int, int]:
    prs = Presentation(str(input_path))
    removed = 0
    scanned = 0
    for slide in prs.slides:
        for shape in list(slide.shapes):
            scanned += 1
            if should_remove(shape, remove_all_placeholders):
                remove_shape(shape)
                removed += 1
    prs.save(str(output_path))
    return scanned, removed


def main() -> None:
    parser = argparse.ArgumentParser(description="Remove unused template placeholder text boxes from PPTX files.")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--remove-all-placeholders",
        action="store_true",
        help="Remove placeholders even when they contain non-prompt text. Use carefully.",
    )
    args = parser.parse_args()

    scanned, removed = cleanup(args.input, args.output, args.remove_all_placeholders)
    print(f"scanned_shapes={scanned}")
    print(f"removed_shapes={removed}")
    print(f"output={args.output}")


if __name__ == "__main__":
    main()
