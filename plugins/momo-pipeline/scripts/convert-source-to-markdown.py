#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import html
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".log"}
CSV_EXTENSIONS = {".csv", ".tsv"}


def read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp950", "big5", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(errors="replace")


def write_markdown(out: Path, source: Path, method: str, body: str) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    text = body.strip()
    if not text:
        text = "_No extractable text was found._"
    out.write_text(
        f"# {source.name}\n\n"
        f"- Source: `{source}`\n"
        f"- Conversion method: `{method}`\n\n"
        f"{text}\n",
        encoding="utf-8",
    )


def try_markitdown(source: Path, verbose: bool = False) -> str | None:
    try:
        from markitdown import MarkItDown  # type: ignore
    except Exception as exc:
        if verbose:
            print(f"Plugin MarkItDown converter unavailable; using built-in fallback if possible: {exc}", file=sys.stderr)
        return None

    try:
        result = MarkItDown().convert(str(source))
        text = getattr(result, "text_content", None)
        if text and str(text).strip():
            return str(text)
    except Exception as exc:
        if verbose:
            print(f"MarkItDown failed; using built-in fallback if possible: {exc}", file=sys.stderr)
    return None


def pptx_to_markdown(source: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore
    except Exception as exc:
        raise RuntimeError(f"PPTX fallback requires python-pptx: {exc}") from exc

    prs = Presentation(str(source))
    lines = [
        f"- Slides: {len(prs.slides)}",
        f"- Size: {round(prs.slide_width / 914400, 3)} x {round(prs.slide_height / 914400, 3)} in",
        "",
    ]

    for slide_number, slide in enumerate(prs.slides, start=1):
        image_count = 0
        table_count = 0
        text_blocks: list[str] = []
        tables: list[list[list[str]]] = []

        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    text_blocks.append(text)
            if getattr(shape, "has_table", False):
                table_count += 1
                rows: list[list[str]] = []
                for row in shape.table.rows:
                    rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
                tables.append(rows)

        lines.append(f"## Slide {slide_number}: {slide.slide_layout.name}")
        lines.append("")
        lines.append(f"- Images: {image_count}")
        lines.append(f"- Tables: {table_count}")
        lines.append("")

        for block in text_blocks:
            lines.append(block)
            lines.append("")

        for rows in tables:
            if not rows:
                continue
            max_cols = max(len(row) for row in rows)
            padded = [row + [""] * (max_cols - len(row)) for row in rows]
            lines.append("| " + " | ".join(padded[0]) + " |")
            lines.append("| " + " | ".join(["---"] * max_cols) + " |")
            for row in padded[1:]:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")

    return "\n".join(lines)


def docx_to_markdown(source: Path) -> str:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(source) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    lines: list[str] = []

    for child in root.findall(".//w:body/*", ns):
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            texts = [node.text or "" for node in child.findall(".//w:t", ns)]
            para = "".join(texts).strip()
            if para:
                lines.append(para)
                lines.append("")
        elif tag == "tbl":
            rows: list[list[str]] = []
            for tr in child.findall(".//w:tr", ns):
                row: list[str] = []
                for tc in tr.findall("./w:tc", ns):
                    cell_text = "".join(node.text or "" for node in tc.findall(".//w:t", ns)).strip()
                    row.append(cell_text)
                rows.append(row)
            if rows:
                max_cols = max(len(row) for row in rows)
                padded = [row + [""] * (max_cols - len(row)) for row in rows]
                lines.append("| " + " | ".join(padded[0]) + " |")
                lines.append("| " + " | ".join(["---"] * max_cols) + " |")
                for row in padded[1:]:
                    lines.append("| " + " | ".join(row) + " |")
                lines.append("")

    return "\n".join(lines)


def csv_to_markdown(source: Path) -> str:
    delimiter = "\t" if source.suffix.lower() == ".tsv" else ","
    text = read_text(source)
    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    if not rows:
        return ""
    max_cols = max(len(row) for row in rows)
    padded = [row + [""] * (max_cols - len(row)) for row in rows]
    lines = ["| " + " | ".join(padded[0]) + " |", "| " + " | ".join(["---"] * max_cols) + " |"]
    for row in padded[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def html_to_markdown(source: Path) -> str:
    text = read_text(source)
    text = re.sub(r"(?is)<script.*?</script>", " ", text)
    text = re.sub(r"(?is)<style.*?</style>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = re.sub(r"(?i)</p\s*>", "\n\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def fallback_convert(source: Path) -> tuple[str, str]:
    ext = source.suffix.lower()
    if ext == ".pptx":
        return "built-in-pptx", pptx_to_markdown(source)
    if ext == ".docx":
        return "built-in-docx", docx_to_markdown(source)
    if ext in TEXT_EXTENSIONS:
        return "built-in-text", read_text(source)
    if ext in CSV_EXTENSIONS:
        return "built-in-csv", csv_to_markdown(source)
    if ext in {".html", ".htm"}:
        return "built-in-html", html_to_markdown(source)
    raise RuntimeError(
        f"No built-in fallback for {ext or 'files without extension'}. "
        "Install MarkItDown in the plugin environment or inspect the source directly."
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert source documents to Markdown without relying on a global MarkItDown command.")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--no-markitdown", action="store_true", help="Skip MarkItDown and use built-in fallback directly.")
    parser.add_argument("--verbose", action="store_true", help="Print MarkItDown fallback diagnostics.")
    args = parser.parse_args()

    source = args.source.expanduser().resolve()
    output = args.output.expanduser().resolve()
    if not source.exists():
        raise SystemExit(f"Source file not found: {source}")

    if not args.no_markitdown:
        text = try_markitdown(source, verbose=args.verbose)
        if text:
            write_markdown(output, source, "plugin-markitdown", text)
            print(f"converted={output}")
            print("method=plugin-markitdown")
            return

    method, body = fallback_convert(source)
    write_markdown(output, source, method, body)
    print(f"converted={output}")
    print(f"method={method}")


if __name__ == "__main__":
    main()
